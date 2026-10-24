import csv
import hashlib
import json
import logging
import math
import sqlite3
import tempfile
from collections import Counter
from collections.abc import Iterator
from contextlib import contextmanager
from datetime import UTC, date, datetime
from itertools import chain, islice
from pathlib import Path
from time import monotonic
from typing import Any
from uuid import UUID

import ijson
import openpyxl
import pyarrow as pa
import xlrd
from pyarrow import parquet
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.crud import crud_data_quality, crud_dataset
from app.database.database import SessionLocal
from app.models.data_quality import DataQualityCheck
from app.models.dataset import DatasetColumnMapping, DatasetVersion
from app.models.user import User
from app.schemas.data_quality import (
    DataQualityCheckStarted,
    DataQualityCheckStatus,
    DataQualityColumnCount,
    DataQualityIssueCount,
    DataQualityIssueItem,
    DataQualityIssuesPage,
    DataQualityOutlierRule,
    DataQualityReport,
)
from app.services.dataset_preview_service import (
    JSON_ARRAY_FIELDS,
    PNEUMA_HEADERS,
    PNEUMA_POINT_FIELDS,
    _resolve_file_path,
)

logger = logging.getLogger(__name__)
MAX_COLUMNS = 512
SAMPLE_ROWS = 1000
DB_BATCH_SIZE = 1000
PROGRESS_BATCH_SIZE = 5000
PROGRESS_INTERVAL_SECONDS = 2.0
PROGRESS_POINT_BATCH_SIZE = 1024
NULL_TOKENS = {"", "null", "none", "nan"}
OUTLIER_FIELDS = {"speed", "avg_speed", "lon_acc", "lat_acc"}
ALIASES = {
    "track_id": {"track_id", "vehicle_id", "trackid", "vehicleid"},
    "timestamp": {"timestamp", "time", "timestamp_ms"},
    "latitude": {"latitude", "lat"},
    "longitude": {"longitude", "lon", "lng"},
}


class DataQualityNotFoundError(FileNotFoundError):
    """Raised when a dataset version or quality check cannot be found."""


class DataQualityConflictError(ValueError):
    """Raised when a quality check is already queued or running."""


class DataQualitySchemaError(ValueError):
    """Raised when a dataset has no usable tabular schema."""


class DataQualityReportNotReadyError(ValueError):
    """Raised when a report is requested before a check completes."""


def _resolve_dataset_version(
    db: Session,
    dataset_id: UUID,
) -> tuple[Any, DatasetVersion]:
    result = crud_dataset.get_dataset_latest_version(db, dataset_id)
    if result is None:
        raise DataQualityNotFoundError("Không tìm thấy dataset hoặc phiên bản đã tải lên")
    return result


def _value_matches_type(value: Any, expected_type: str) -> bool:
    if expected_type == "unknown" or expected_type == "string":
        return True
    if expected_type == "number":
        try:
            return math.isfinite(float(value))
        except (TypeError, ValueError):
            return False
    if expected_type == "integer":
        if isinstance(value, bool):
            return False
        try:
            return float(value).is_integer()
        except (TypeError, ValueError):
            return False
    if expected_type == "float":
        try:
            return math.isfinite(float(value))
        except (TypeError, ValueError):
            return False
    if expected_type == "boolean":
        return isinstance(value, bool) or str(value).strip().lower() in {
            "true",
            "false",
        }
    if expected_type == "date":
        if isinstance(value, (date, datetime)):
            return True
        try:
            date.fromisoformat(str(value).strip())
            return True
        except ValueError:
            return False
    if expected_type == "datetime":
        if isinstance(value, datetime):
            return True
        try:
            datetime.fromisoformat(str(value).strip())
            return True
        except ValueError:
            return False
    return True


def _get_mappings(
    db: Session,
    dataset_id: UUID,
    version_id: UUID,
) -> dict[str, str]:
    record: DatasetColumnMapping | None = crud_dataset.get_dataset_column_mapping(
        db,
        dataset_id,
        version_id,
    )
    if record is None:
        return {}
    return {
        str(item["source_column"]): str(item["target_column"])
        for item in record.mappings
        if item.get("target_column") is not None
    }


def _version_path(version: DatasetVersion) -> Path:
    try:
        return _resolve_file_path(version)
    except FileNotFoundError as exc:
        raise DataQualityNotFoundError(str(exc)) from exc


def _status_schema(check: DataQualityCheck) -> DataQualityCheckStatus:
    return DataQualityCheckStatus(
        check_id=check.id,
        dataset_id=check.dataset_id,
        dataset_version_id=check.dataset_version_id,
        status=check.status,
        progress_percent=check.progress_percent,
        processed_rows=check.processed_rows,
        total_rows=check.total_rows,
        error_message=check.error_message,
        created_at=check.created_at,
        started_at=check.started_at,
        completed_at=check.completed_at,
    )


def start_data_quality_check(
    db: Session,
    *,
    dataset_id: UUID,
    user: User,
) -> tuple[DataQualityCheckStarted, Path]:
    _, version = _resolve_dataset_version(db, dataset_id)
    path = _version_path(version)
    active = crud_data_quality.get_active_check(db, version.id)
    if active is not None:
        raise DataQualityConflictError(
            f"Dataset đang được kiểm tra (check_id={active.id})"
        )
    check = DataQualityCheck(
        dataset_id=dataset_id,
        dataset_version_id=version.id,
        started_by=user.id,
        status="queued",
        progress_percent=0,
        processed_rows=0,
    )
    db.add(check)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise DataQualityConflictError(
            "Dataset đang có một tác vụ kiểm tra đang chờ hoặc đang chạy"
        ) from exc
    db.refresh(check)
    response = DataQualityCheckStarted(
        **_status_schema(check).model_dump(),
        message="Đã xếp hàng kiểm tra chất lượng dataset",
    )
    return response, path


def get_data_quality_status(
    db: Session,
    *,
    dataset_id: UUID,
    check_id: UUID,
) -> DataQualityCheckStatus:
    check = crud_data_quality.get_check(db, dataset_id, check_id)
    if check is None:
        raise DataQualityNotFoundError("Không tìm thấy tác vụ kiểm tra chất lượng")
    return _status_schema(check)


def get_data_quality_report(
    db: Session,
    *,
    dataset_id: UUID,
    check_id: UUID,
) -> DataQualityReport:
    check = crud_data_quality.get_check(db, dataset_id, check_id)
    if check is None:
        raise DataQualityNotFoundError("Không tìm thấy tác vụ kiểm tra chất lượng")
    if check.status != "completed" or check.report is None:
        if check.status == "failed":
            raise DataQualityReportNotReadyError(
                f"Tác vụ kiểm tra thất bại: {check.error_message or 'lỗi không xác định'}"
            )
        raise DataQualityReportNotReadyError(
            f"Báo cáo chưa sẵn sàng (status={check.status})"
        )
    return DataQualityReport(check_id=check.id, **check.report)


def get_data_quality_issues(
    db: Session,
    *,
    dataset_id: UUID,
    check_id: UUID,
    page: int,
    page_size: int,
) -> DataQualityIssuesPage:
    check = crud_data_quality.get_check(db, dataset_id, check_id)
    if check is None:
        raise DataQualityNotFoundError("Không tìm thấy tác vụ kiểm tra chất lượng")
    issues, total = crud_data_quality.list_issues(
        db,
        check_id,
        offset=(page - 1) * page_size,
        limit=page_size,
    )
    return DataQualityIssuesPage(
        check_id=check_id,
        page=page,
        page_size=page_size,
        total=total,
        issues=[
            DataQualityIssueItem(
                row_number=issue.row_number,
                issue_type=issue.issue_type,
                severity="warning" if issue.issue_type == "outlier" else "error",
                column_name=issue.column_name,
                message=issue.message,
                observed_value=issue.observed_value,
            )
            for issue in issues
        ],
    )


def _clean_headers(values: Iterator[Any]) -> tuple[list[str], list[str]]:
    headers: list[str] = []
    issues: list[str] = []
    seen: Counter[str] = Counter()
    for index, raw_value in enumerate(values, start=1):
        name = str(raw_value).strip() if raw_value is not None else ""
        if not name:
            issues.append(f"Header cột {index} bị trống")
            name = f"column_{index}"
        seen[name] += 1
        if seen[name] > 1:
            issues.append(f"Header '{name}' bị trùng")
            name = f"{name}_{seen[name]}"
        headers.append(name)
    if len(headers) > MAX_COLUMNS:
        raise DataQualitySchemaError(f"Dataset vượt quá giới hạn {MAX_COLUMNS} cột")
    if not headers:
        raise DataQualitySchemaError("Dataset không có header hợp lệ")
    return headers, issues


def _json_array_prefix(path: Path) -> str | None:
    try:
        with path.open("rb") as source:
            for prefix, event, _ in ijson.parse(source):
                if event == "start_array" and (
                    prefix == "" or prefix in JSON_ARRAY_FIELDS
                ):
                    return prefix
    except (OSError, ijson.JSONError) as exc:
        raise DataQualitySchemaError("Không đọc được cấu trúc JSON") from exc
    return None


def _json_record_values(record: Any, headers: list[str] | None = None) -> list[Any]:
    if isinstance(record, dict):
        return [record.get(name) for name in headers] if headers else list(record.values())
    if isinstance(record, (list, tuple)):
        return list(record)
    return [record]


@contextmanager
def _open_records(
    path: Path,
    file_format: str,
) -> Iterator[
    tuple[
        list[str],
        list[str],
        str,
        Iterator[tuple[int, list[Any], str | None]],
        int | None,
    ]
]:
    suffix = path.suffix.lower()
    format_name = file_format.lower()
    if suffix == ".csv" or format_name == "csv":
        try:
            source = path.open("r", encoding="utf-8-sig", newline="")
            first_line = source.readline(64 * 1024)
            source.seek(0)
            try:
                dialect = csv.Sniffer().sniff(first_line, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel()
                dialect.delimiter = max((",", ";", "\t"), key=first_line.count)
                if not first_line.count(dialect.delimiter):
                    dialect.delimiter = ","
            reader = csv.reader(source, dialect, strict=True)
            try:
                raw_headers = next(reader)
            except StopIteration as exc:
                source.close()
                raise DataQualitySchemaError("CSV rỗng, không có header") from exc
            headers, header_issues = _clean_headers(iter(raw_headers))
            is_pneuma = tuple(name.lower() for name in headers) == PNEUMA_HEADERS

            def csv_records() -> Iterator[tuple[int, list[Any], str | None]]:
                while True:
                    try:
                        values = next(reader)
                    except StopIteration:
                        return
                    except csv.Error as exc:
                        yield reader.line_num, [], str(exc)
                        continue
                    yield reader.line_num, values, None

            try:
                yield (
                    headers,
                    header_issues,
                    "pneuma_wide_trajectory" if is_pneuma else "tabular_csv",
                    csv_records(),
                    None,
                )
            finally:
                source.close()
            return
        except OSError as exc:
            raise DataQualitySchemaError("Không mở được file CSV") from exc

    if suffix == ".parquet" or format_name == "parquet":
        try:
            parquet_file = parquet.ParquetFile(path)
            schema = parquet_file.schema_arrow
            headers, header_issues = _clean_headers(iter(schema.names))

            def parquet_records() -> Iterator[tuple[int, list[Any], str | None]]:
                row_number = 2
                for batch in parquet_file.iter_batches(batch_size=2048):
                    for offset in range(batch.num_rows):
                        yield (
                            row_number,
                            [
                                batch.column(index)[offset].as_py()
                                for index in range(batch.num_columns)
                            ],
                            None,
                        )
                        row_number += 1

            yield headers, header_issues, "tabular_parquet", parquet_records(), parquet_file.metadata.num_rows
            return
        except (OSError, pa.ArrowException, ValueError) as exc:
            raise DataQualitySchemaError("Không đọc được schema Parquet") from exc

    if suffix == ".json" or format_name == "json":
        prefix = _json_array_prefix(path)
        if prefix is None:
            raise DataQualitySchemaError(
                "JSON cần là mảng records hoặc có mảng records/data/rows/results"
            )
        item_prefix = f"{prefix}.item" if prefix else "item"
        try:
            source = path.open("rb")
            iterator = iter(ijson.items(source, item_prefix))
            first_record = next(iterator, None)
            if first_record is None:
                source.close()
                raise DataQualitySchemaError("JSON không chứa records")
            if isinstance(first_record, dict):
                headers, header_issues = _clean_headers(iter(first_record.keys()))
            elif isinstance(first_record, (list, tuple)):
                headers, header_issues = _clean_headers(
                    iter(f"column_{index}" for index in range(1, len(first_record) + 1))
                )
            else:
                headers, header_issues = ["value"], []

            def json_records() -> Iterator[tuple[int, list[Any], str | None]]:
                yield 2, _json_record_values(first_record, headers), None
                for row_number, record in enumerate(iterator, start=3):
                    if isinstance(record, dict):
                        values = [record.get(name) for name in headers]
                        extra_keys = set(record) - set(headers)
                        missing_keys = set(headers) - set(record)
                        errors = []
                        if extra_keys:
                            extra_values = {
                                str(key): record[key]
                                for key in sorted(extra_keys, key=str)
                            }
                            values.append(
                                "__extra_fields__="
                                + json.dumps(extra_values, ensure_ascii=False, default=str)
                            )
                            errors.append(
                                "trường ngoài schema: "
                                + ", ".join(sorted(str(key) for key in extra_keys)[:10])
                            )
                        if missing_keys:
                            errors.append(
                                "thiếu trường: "
                                + ", ".join(sorted(str(key) for key in missing_keys)[:10])
                            )
                        parse_error = (
                            "JSON record " + "; ".join(errors) if errors else None
                        )
                        yield row_number, values, parse_error
                    else:
                        values = _json_record_values(record, headers)
                        parse_error = (
                            f"JSON record có {len(values)} trường, schema có {len(headers)}"
                            if len(values) != len(headers)
                            else None
                        )
                        yield row_number, values, parse_error

            try:
                yield headers, header_issues, "json_records", json_records(), None
            finally:
                source.close()
            return
        except (OSError, ijson.JSONError) as exc:
            raise DataQualitySchemaError("JSON records không hợp lệ") from exc

    if suffix == ".xlsx" or format_name == "xlsx":
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            worksheet = workbook.active
            iterator = worksheet.iter_rows(values_only=True)
            raw_headers = next(iterator, ())
            headers, header_issues = _clean_headers(iter(raw_headers))

            def xlsx_records() -> Iterator[tuple[int, list[Any], str | None]]:
                for row_number, values in enumerate(iterator, start=2):
                    yield row_number, list(values), None

            try:
                yield headers, header_issues, "tabular_xlsx", xlsx_records(), worksheet.max_row - 1
            finally:
                workbook.close()
            return
        except (OSError, ValueError, StopIteration) as exc:
            raise DataQualitySchemaError("Không đọc được workbook XLSX") from exc

    if suffix == ".xls" or format_name == "xls":
        try:
            workbook = xlrd.open_workbook(path, on_demand=True)
            if workbook.nsheets == 0:
                raise DataQualitySchemaError("Excel không có worksheet")
            worksheet = workbook.sheet_by_index(0)
            headers, header_issues = _clean_headers(
                iter(worksheet.row_values(0) if worksheet.nrows else [])
            )

            def xls_records() -> Iterator[tuple[int, list[Any], str | None]]:
                for row_index in range(1, worksheet.nrows):
                    yield row_index + 1, worksheet.row_values(row_index), None

            try:
                yield headers, header_issues, "tabular_xls", xls_records(), worksheet.nrows - 1
            finally:
                workbook.release_resources()
            return
        except (OSError, xlrd.XLRDError, ValueError) as exc:
            raise DataQualitySchemaError("Không đọc được workbook XLS") from exc

    raise DataQualitySchemaError(
        f"Data Quality Check chưa hỗ trợ định dạng '{file_format}'"
    )


def _is_missing(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return isinstance(value, str) and value.strip().lower() in NULL_TOKENS


def _value_kind(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "integer" if value.is_integer() else "float"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    text_value = str(value).strip()
    if text_value.lower() in {"true", "false"}:
        return "boolean"
    try:
        int(text_value)
        return "integer"
    except ValueError:
        pass
    try:
        number = float(text_value)
        return "float" if math.isfinite(number) else "invalid"
    except (ValueError, TypeError):
        pass
    try:
        date.fromisoformat(text_value)
        return "date"
    except ValueError:
        pass
    try:
        datetime.fromisoformat(text_value)
        return "datetime"
    except ValueError:
        return "string"


def _infer_column_types(
    headers: list[str],
    sample: list[list[Any]],
    mappings: dict[str, str],
    is_pneuma: bool,
) -> dict[str, str]:
    result: dict[str, str] = {}
    numeric_names = {
        name.lower()
        for name in ("latitude", "lat", "longitude", "lon", "lng", "speed", "average_speed",
                     "avg_speed", "x", "y", "timestamp", "time", "timestamp_ms")
    }
    for index, header in enumerate(headers):
        logical_name = mappings.get(header, header).lower()
        if is_pneuma or logical_name in numeric_names:
            result[header] = "number"
            continue
        kinds = Counter(
            _value_kind(row[index])
            for row in sample
            if index < len(row) and not _is_missing(row[index])
        )
        if not kinds:
            result[header] = "unknown"
            continue
        kind, count = kinds.most_common(1)[0]
        result[header] = (
            kind if kind != "invalid" and count / sum(kinds.values()) >= 0.7 else "unknown"
        )
    return result


def _parse_timestamp(value: Any) -> bool:
    if isinstance(value, bool) or value is None:
        return False
    if isinstance(value, (int, float)):
        return math.isfinite(float(value)) and value >= 0
    text_value = str(value).strip()
    try:
        number = float(text_value)
        return math.isfinite(number) and number >= 0
    except ValueError:
        pass
    try:
        datetime.fromisoformat(text_value)
        return True
    except ValueError:
        try:
            date.fromisoformat(text_value)
            return True
        except ValueError:
            return False


def _parse_coordinate(value: Any, *, latitude: bool) -> tuple[bool, bool]:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return False, False
    if not math.isfinite(number):
        return False, False
    low, high = (-90, 90) if latitude else (-180, 180)
    return True, low <= number <= high


def _normalize_pneuma_trailing_delimiter(values: list[Any]) -> list[Any]:
    if (
        values
        and _is_missing(values[-1])
        and (len(values) - len(PNEUMA_HEADERS)) % len(PNEUMA_POINT_FIELDS) == 1
    ):
        return values[:-1]
    return values


def _string_value(value: Any) -> str:
    text_value = str(value)
    return text_value if len(text_value) <= 256 else text_value[:253] + "..."


def _semantic_sources(
    headers: list[str],
    mappings: dict[str, str],
) -> dict[str, str | None]:
    result: dict[str, str | None] = {}
    for semantic, aliases in ALIASES.items():
        result[semantic] = next(
            (
                source
                for source in headers
                if source.lower() in aliases
                or mappings.get(source, "").lower() in aliases
            ),
            None,
        )
    return result


def _run_check(
    check_id: UUID,
    path: Path,
    version: DatasetVersion,
    mappings: dict[str, str],
    *,
    outlier_rules: list[DataQualityOutlierRule] | None = None,
) -> None:
    db = SessionLocal()
    temp_dir = tempfile.TemporaryDirectory(prefix="traffic-quality-")
    disk_state = sqlite3.connect(str(Path(temp_dir.name) / "quality_state.sqlite"))
    disk_state.execute("CREATE TABLE seen_hashes (digest BLOB PRIMARY KEY)")
    disk_state.execute("CREATE TABLE affected_rows (row_number INTEGER PRIMARY KEY)")
    disk_state.execute(
        "CREATE TABLE issue_affected_rows "
        "(issue_type TEXT NOT NULL, row_number INTEGER NOT NULL, "
        "PRIMARY KEY(issue_type, row_number))"
    )
    try:
        check = db.query(DataQualityCheck).filter(DataQualityCheck.id == check_id).one()
        check.status = "running"
        check.progress_percent = 1
        check.started_at = datetime.now(UTC)
        db.commit()

        error_totals: Counter[str] = Counter()
        column_errors: Counter[str] = Counter()
        warning_totals: Counter[str] = Counter()
        column_warnings: Counter[str] = Counter()
        rules_by_field = {rule.field: rule for rule in outlier_rules or []}
        missing_by_column: Counter[str] = Counter()
        total_values_by_column: Counter[str] = Counter()
        missing_required: list[str] = []
        schema_invalid = False
        total_rows = 0
        processed = 0
        pending_issues: list[dict[str, Any]] = []
        row_issue_types: set[str] = set()
        row_has_error = False
        last_progress_at = monotonic()

        with _open_records(path, version.file_format) as (
            headers,
            header_issues,
            structure,
            records,
            known_total,
        ):
            is_pneuma = structure == "pneuma_wide_trajectory"
            schema_invalid = bool(header_issues)
            logical_sources = _semantic_sources(headers, mappings)
            missing_required = [
                required
                for required in ("track_id", "timestamp")
                if logical_sources[required] is None
            ]
            schema_invalid = schema_invalid or bool(missing_required)
            # pNEUMA has a fixed numeric schema. Sampling 1,000 entire trajectories
            # delays all validation and can retain millions of point values in RAM.
            sample = [] if is_pneuma else list(islice(records, SAMPLE_ROWS))
            sample_values = [values for _, values, parse_error in sample if parse_error is None]
            expected_types = _infer_column_types(
                headers,
                sample_values,
                mappings,
                is_pneuma,
            )

            def flush_issues() -> None:
                crud_data_quality.insert_issues(db, pending_issues)
                pending_issues.clear()

            def publish_progress(*, force: bool = False) -> None:
                nonlocal last_progress_at
                now = monotonic()
                if not force and now - last_progress_at < PROGRESS_INTERVAL_SECONDS:
                    return
                flush_issues()
                current_check = db.query(DataQualityCheck).filter(
                    DataQualityCheck.id == check_id
                ).one()
                current_check.processed_rows = processed
                current_check.total_rows = known_total
                current_check.progress_percent = (
                    max(1, min(95, processed * 95 // known_total))
                    if known_total
                    else min(95, 5 + processed // PROGRESS_BATCH_SIZE)
                )
                db.commit()
                disk_state.commit()
                last_progress_at = now

            def add_issue(
                row_number: int | None,
                issue_type: str,
                column_name: str | None,
                message: str,
                observed_value: Any = None,
            ) -> None:
                nonlocal row_has_error
                pending_issues.append({
                    "check_id": check_id,
                    "row_number": row_number,
                    "issue_type": issue_type,
                    "column_name": column_name,
                    "message": message[:512],
                    "observed_value": (
                        None if observed_value is None else _string_value(observed_value)
                    ),
                })
                is_warning = issue_type == "outlier"
                totals = warning_totals if is_warning else error_totals
                columns = column_warnings if is_warning else column_errors
                totals[issue_type] += 1
                if column_name:
                    columns[column_name] += 1
                if row_number is not None:
                    if not is_warning and not row_has_error:
                        disk_state.execute(
                            "INSERT OR IGNORE INTO affected_rows(row_number) VALUES (?)",
                            (row_number,),
                        )
                        row_has_error = True
                    if issue_type not in row_issue_types:
                        disk_state.execute(
                            "INSERT OR IGNORE INTO issue_affected_rows "
                            "(issue_type, row_number) VALUES (?, ?)",
                            (issue_type, row_number),
                        )
                        row_issue_types.add(issue_type)
                if len(pending_issues) >= DB_BATCH_SIZE:
                    flush_issues()
                    db.commit()
                    disk_state.commit()

            def check_outlier(
                row_number: int,
                column_name: str,
                field_name: str,
                value: Any,
                *,
                number: float | None = None,
            ) -> None:
                rule = rules_by_field.get(field_name)
                if rule is None or _is_missing(value) or isinstance(value, bool):
                    return
                if number is None:
                    try:
                        number = float(value)
                    except (TypeError, ValueError, OverflowError):
                        return
                if not math.isfinite(number):
                    return
                # Negative speeds remain errors, not duplicate outlier warnings.
                if field_name in {"speed", "avg_speed"} and number < 0:
                    return
                if (
                    rule.min_value is not None and number < rule.min_value
                    or rule.max_value is not None and number > rule.max_value
                ):
                    add_issue(
                        row_number,
                        "outlier",
                        column_name,
                        f"{field_name} ngoài ngưỡng cấu hình "
                        f"[{rule.min_value}, {rule.max_value}] {rule.unit}",
                        value,
                    )

            for header_issue in header_issues:
                add_issue(None, "invalid_format", None, header_issue)
            for required in missing_required:
                add_issue(
                    None,
                    "invalid_format",
                    None,
                    f"Thiếu cột bắt buộc '{required}' trong schema",
                )

            timestamp_source = logical_sources["timestamp"]
            latitude_source = logical_sources["latitude"]
            longitude_source = logical_sources["longitude"]
            all_records = chain(sample, records)
            for row_number, raw_values, parse_error in all_records:
                total_rows += 1
                row_issue_types.clear()
                row_has_error = False
                if parse_error is not None:
                    add_issue(
                        row_number,
                        "invalid_format",
                        None,
                        "Dòng không khớp cú pháp hoặc schema của file",
                        parse_error,
                    )
                    if not raw_values:
                        processed += 1
                        publish_progress(force=processed % PROGRESS_BATCH_SIZE == 0)
                        continue

                values = list(raw_values)
                if is_pneuma:
                    values = _normalize_pneuma_trailing_delimiter(values)
                    if (
                        len(values) < len(PNEUMA_HEADERS)
                        or (len(values) - len(PNEUMA_HEADERS)) % len(PNEUMA_POINT_FIELDS)
                    ):
                        add_issue(
                            row_number,
                            "invalid_format",
                            None,
                            "Số trường không khớp cấu trúc pNEUMA wide trajectory",
                            len(values),
                        )
                    digest_values = values
                else:
                    if len(values) != len(headers) and parse_error is None:
                        add_issue(
                            row_number,
                            "invalid_format",
                            None,
                            f"Dòng có {len(values)} trường, header có {len(headers)}",
                            len(values),
                        )
                    digest_values = values

                digest = hashlib.sha256(
                    json.dumps(
                        digest_values,
                        ensure_ascii=False,
                        separators=(",", ":"),
                        default=str,
                    ).encode("utf-8")
                ).digest()
                inserted = disk_state.execute(
                    "INSERT OR IGNORE INTO seen_hashes(digest) VALUES (?)",
                    (digest,),
                ).rowcount
                if not inserted:
                    add_issue(
                        row_number,
                        "duplicate_record",
                        None,
                        "Bản ghi trùng hoàn toàn với một dòng trước đó",
                    )

                if is_pneuma:
                    # Mirror dataset_preview_service: floor division starting from
                    # len(PNEUMA_HEADERS)=10 so the first point group (indices 4-9)
                    # is counted as point 1, additional groups start at index 10.
                    point_count = (
                        1 + (len(values) - len(PNEUMA_HEADERS)) // len(PNEUMA_POINT_FIELDS)
                        if len(values) >= len(PNEUMA_HEADERS)
                        else 0
                    )
                    for field_index, field_name in enumerate(PNEUMA_HEADERS[:4]):
                        value = values[field_index] if field_index < len(values) else None
                        total_values_by_column[field_name] += 1
                        if _is_missing(value):
                            missing_by_column[field_name] += 1
                            add_issue(
                                row_number,
                                "missing_value",
                                field_name,
                                "Giá trị bị thiếu",
                                value,
                            )
                        elif field_name in {"track_id", "traveled_d", "avg_speed"}:
                            try:
                                number = float(value)
                                valid_number = math.isfinite(number)
                                if field_name == "track_id":
                                    valid_number = valid_number and number.is_integer()
                            except (TypeError, ValueError):
                                valid_number = False
                            if not valid_number:
                                add_issue(
                                    row_number,
                                    "invalid_format",
                                    field_name,
                                    "Giá trị không khớp kiểu dữ liệu pNEUMA",
                                    value,
                                )
                            elif field_name == "avg_speed" and number < 0:
                                add_issue(
                                    row_number,
                                    "negative_speed",
                                    field_name,
                                    "Tốc độ phải >= 0",
                                    value,
                                )
                            if valid_number and field_name in rules_by_field:
                                check_outlier(
                                    row_number, field_name, field_name, value, number=number
                                )
                    point_start = 4
                    for point_index in range(point_count):
                        if point_index and point_index % PROGRESS_POINT_BATCH_SIZE == 0:
                            publish_progress()
                        start = point_start + point_index * len(PNEUMA_POINT_FIELDS)
                        point_values = values[start : start + len(PNEUMA_POINT_FIELDS)]
                        if len(point_values) < len(PNEUMA_POINT_FIELDS):
                            add_issue(
                                row_number,
                                "invalid_format",
                                None,
                                "Nhóm điểm quỹ đạo pNEUMA không đủ 6 trường",
                                len(point_values),
                            )
                        for field_offset, field_name in enumerate(PNEUMA_POINT_FIELDS):
                            value = (
                                point_values[field_offset]
                                if field_offset < len(point_values)
                                else None
                            )
                            total_values_by_column[field_name] += 1
                            if _is_missing(value):
                                missing_by_column[field_name] += 1
                                add_issue(
                                    row_number,
                                    "missing_value",
                                    field_name,
                                    "Giá trị điểm quỹ đạo bị thiếu",
                                    value,
                                )
                                continue
                            try:
                                number = float(value)
                                valid_number = math.isfinite(number)
                            except (TypeError, ValueError):
                                valid_number = False
                                number = 0
                            if not valid_number:
                                issue_type = (
                                    "invalid_timestamp"
                                    if field_name == "time"
                                    else "invalid_coordinate"
                                    if field_name in {"lat", "lon"}
                                    else "invalid_format"
                                )
                                add_issue(
                                    row_number,
                                    issue_type,
                                    field_name,
                                    "Giá trị không phải số hữu hạn",
                                    value,
                                )
                            elif field_name == "speed" and number < 0:
                                add_issue(
                                    row_number,
                                    "negative_speed",
                                    field_name,
                                    "Tốc độ phải >= 0",
                                    value,
                                )
                            elif field_name == "time" and number < 0:
                                add_issue(
                                    row_number,
                                    "invalid_timestamp",
                                    field_name,
                                    "Thời gian pNEUMA tương đối phải >= 0 giây",
                                    value,
                                )
                            elif field_name in {"lat", "lon"}:
                                limit = 90 if field_name == "lat" else 180
                                if not -limit <= number <= limit:
                                    add_issue(
                                        row_number,
                                        "invalid_coordinate",
                                        field_name,
                                        f"Giá trị nằm ngoài miền [-{limit}, {limit}]",
                                        value,
                                    )
                            if valid_number and field_name in rules_by_field:
                                check_outlier(
                                    row_number, field_name, field_name, value, number=number
                                )
                else:
                    row = {
                        header: values[index] if index < len(values) else None
                        for index, header in enumerate(headers)
                    }
                    for header in headers:
                        value = row[header]
                        total_values_by_column[header] += 1
                        if _is_missing(value):
                            missing_by_column[header] += 1
                            add_issue(
                                row_number,
                                "missing_value",
                                header,
                                "Giá trị bị thiếu",
                                value,
                            )
                            continue

                        logical_name = mappings.get(header, header).lower()
                        if header == timestamp_source and not _parse_timestamp(value):
                            add_issue(
                                row_number,
                                "invalid_timestamp",
                                header,
                                "Timestamp không đúng định dạng hoặc miền giá trị",
                                value,
                            )
                        elif header == latitude_source or logical_name in ALIASES["latitude"]:
                            numeric, in_bounds = _parse_coordinate(value, latitude=True)
                            if not numeric or not in_bounds:
                                add_issue(
                                    row_number,
                                    "invalid_coordinate",
                                    header,
                                    "Latitude phải là số trong [-90, 90]",
                                    value,
                                )
                        elif header == longitude_source or logical_name in ALIASES["longitude"]:
                            numeric, in_bounds = _parse_coordinate(value, latitude=False)
                            if not numeric or not in_bounds:
                                add_issue(
                                    row_number,
                                    "invalid_coordinate",
                                    header,
                                    "Longitude phải là số trong [-180, 180]",
                                    value,
                                )
                        elif (
                            expected_types.get(header) != "unknown"
                            and not _value_matches_type(
                                value,
                                expected_types.get(header, "unknown"),
                            )
                        ):
                            add_issue(
                                row_number,
                                "invalid_format",
                                header,
                                "Giá trị không khớp kiểu dữ liệu suy luận",
                                value,
                            )

                        # Preserve source column names and validate finiteness before sign.
                        if (
                            {header.lower(), logical_name} & {"speed", "avg_speed"}
                            and _value_matches_type(value, "number")
                            and float(value) < 0
                        ):
                            add_issue(
                                row_number,
                                "negative_speed",
                                header,
                                "Tốc độ phải >= 0",
                                value,
                            )

                        # Mapping takes precedence; retain canonical source names as fallback.
                        outlier_field = (
                            logical_name if logical_name in OUTLIER_FIELDS else header.lower()
                        )
                        if outlier_field in rules_by_field:
                            check_outlier(row_number, header, outlier_field, value)

                processed += 1
                publish_progress(force=processed % PROGRESS_BATCH_SIZE == 0)

            flush_issues()
        db.commit()
        disk_state.commit()
        affected_total = disk_state.execute(
            "SELECT COUNT(*) FROM affected_rows"
        ).fetchone()[0]
        errors_by_type: dict[str, DataQualityIssueCount] = {}
        for issue_type, count in error_totals.items():
            type_affected = disk_state.execute(
                "SELECT COUNT(*) FROM issue_affected_rows WHERE issue_type = ?",
                (issue_type,),
            ).fetchone()[0]
            if schema_invalid and issue_type == "invalid_format":
                type_affected = max(type_affected, total_rows)
            errors_by_type[issue_type] = DataQualityIssueCount(
                total_errors=count,
                affected_rows=type_affected,
            )
        if schema_invalid:
            affected_total = max(affected_total, total_rows)

        warning_rows = disk_state.execute(
            "SELECT COUNT(DISTINCT row_number) FROM issue_affected_rows "
            "WHERE issue_type = 'outlier'"
        ).fetchone()[0]

        column_summary = [
            DataQualityColumnCount(
                column_name=name,
                missing_values=missing_by_column[name],
                total_values=total_values_by_column[name],
                missing_percent=round(
                    missing_by_column[name] * 100 / total_values_by_column[name], 2
                )
                if total_values_by_column[name]
                else 0,
                errors=column_errors[name],
            )
            for name in sorted(set(total_values_by_column) | set(column_errors))
        ]
        report = {
            "dataset_id": str(check.dataset_id),
            "dataset_version_id": str(check.dataset_version_id),
            "status": "completed",
            "detected_structure": structure,
            "total_rows": total_rows,
            "valid_rows": max(0, total_rows - affected_total),
            "invalid_rows": min(total_rows, affected_total),
            "total_errors": sum(error_totals.values()),
            "total_warnings": sum(warning_totals.values()),
            "warning_rows": warning_rows,
            "warnings_by_type": {
                "outlier": {
                    "total_warnings": warning_totals["outlier"],
                    "affected_rows": warning_rows,
                }
            } if warning_totals else {},
            "warnings_by_column": dict(column_warnings),
            "outlier_rules": [rule.model_dump() for rule in outlier_rules or []],
            "errors_by_type": {
                issue_type: count.model_dump()
                for issue_type, count in errors_by_type.items()
            },
            "errors_by_column": [item.model_dump() for item in column_summary],
            "missing_values_by_column": [
                item.model_dump()
                for item in column_summary
                if item.missing_values or item.total_values
            ],
            "missing_required_columns": missing_required,
            "generated_at": datetime.now(UTC).isoformat(),
        }
        check = db.query(DataQualityCheck).filter(DataQualityCheck.id == check_id).one()
        check.status = "completed"
        check.progress_percent = 100
        check.processed_rows = total_rows
        check.total_rows = total_rows
        check.report = report
        check.completed_at = datetime.now(UTC)
        db.commit()
    except Exception as exc:
        db.rollback()
        logger.exception(
            "Data quality check failed",
            extra={"check_id": str(check_id), "dataset_version_id": str(version.id)},
        )
        failed = db.query(DataQualityCheck).filter(DataQualityCheck.id == check_id).one_or_none()
        if failed is not None:
            failed.status = "failed"
            failed.error_message = (
                str(exc)[:1000]
                if isinstance(exc, (DataQualitySchemaError, OSError, ValueError))
                else "Tác vụ kiểm tra chất lượng thất bại; xem log server theo check_id"
            )
            failed.completed_at = datetime.now(UTC)
            db.commit()
    finally:
        disk_state.close()
        temp_dir.cleanup()
        db.close()


def run_data_quality_check(
    check_id: UUID,
    *,
    outlier_rules: list[DataQualityOutlierRule] | None = None,
) -> None:
    db = SessionLocal()
    try:
        check = db.query(DataQualityCheck).filter(DataQualityCheck.id == check_id).one_or_none()
        if check is None:
            raise DataQualityNotFoundError("Không tìm thấy tác vụ kiểm tra chất lượng")
        result = crud_dataset.get_dataset_version(
            db,
            check.dataset_id,
            check.dataset_version_id,
        )
        if result is None:
            raise DataQualityNotFoundError("Không tìm thấy phiên bản dataset được kiểm tra")
        _, version = result
        path = _version_path(version)
        mappings = _get_mappings(db, check.dataset_id, check.dataset_version_id)
        _run_check(check_id, path, version, mappings, outlier_rules=outlier_rules)
    except Exception:
        logger.exception("Could not start data quality worker", extra={"check_id": str(check_id)})
        check = db.query(DataQualityCheck).filter(DataQualityCheck.id == check_id).one_or_none()
        if check is not None and check.status in {"queued", "running"}:
            check.status = "failed"
            check.error_message = "Không thể khởi chạy tác vụ kiểm tra chất lượng"
            check.completed_at = datetime.now(UTC)
            db.commit()
    finally:
        db.close()
